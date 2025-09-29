import os
import json
import traceback
from pptx import Presentation

from generate_input_json import GeneratePPTJSON


def generate_json_from_requirements(
    model_to_use="Claude",
    requirements_file="requirements.docx",
    questionnaire_file=None,
    presentation_type="Proposal",
    output_file_name="proposal_output.json",
    max_tokens=4000
):
    """
    Generate presentation JSON from requirements document.
    
    Returns:
        tuple: (success: bool, result: str) - result is either error message or output filename
    """
    try:
        generator = GeneratePPTJSON(
            model_to_use=model_to_use,
            output_file_name=output_file_name,
            max_tokens=max_tokens
        )
        
        message = generator.generate_from_files(
            requirements_file=requirements_file,
            questionnaire_file=questionnaire_file,
            presentation_type=presentation_type
        )
        
        print(f"JSON generation completed: {message}")
        return True, output_file_name
        
    except Exception as e:
        error_msg = f"Error generating JSON: {e}"
        print(error_msg)
        traceback.print_exc()
        return False, error_msg


def add_slide_content(slide, slide_data):
    """
    Add content to a slide based on slide_data structure.
    
    Args:
        slide: PowerPoint slide object
        slide_data: Dictionary containing slide title and content
    """
    # Set title
    if slide.shapes.title:
        slide.shapes.title.text = slide_data.get("title", "")
    
    # Build content text
    content_text = ""
    
    if "content" in slide_data:
        content_text = slide_data["content"]
    elif "sections" in slide_data:
        for section in slide_data["sections"]:
            section_title = section.get("subtitle", section.get("title", ""))
            section_content = section.get("content", "")
            
            if section_title:
                content_text += f"{section_title}\n{section_content}\n\n"
            else:
                content_text += f"{section_content}\n\n"
    
    # Add content to body placeholder if available
    if len(slide.placeholders) > 1:
        body_shape = slide.placeholders[1]
        text_frame = body_shape.text_frame
        text_frame.clear()
        text_frame.text = content_text.strip()


def generate_presentation_from_json(
    json_file_path,
    template_path,
    output_path="Generated_Proposal.pptx",
    slide_layout_index=1
):
    """
    Generate PowerPoint presentation from JSON file.
    
    Args:
        json_file_path: Path to JSON file with slide data
        template_path: Path to PowerPoint template
        output_path: Path for output presentation
        slide_layout_index: Index of slide layout to use from template
    """
    # Verify files exist
    if not os.path.exists(json_file_path):
        raise FileNotFoundError(f"JSON file not found: {json_file_path}")
    
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Template file not found: {template_path}")
    
    # Load JSON data
    print(f"Loading JSON from: {json_file_path}")
    with open(json_file_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    # Load template
    print(f"Loading template from: {template_path}")
    prs = Presentation(template_path)
    
    # Get presentation metadata
    presentation_type = data.get("presentation_type", "Presentation")
    project_title = data.get("project_title", "")
    
    print(f"Generating {presentation_type} presentation with {len(data['slides'])} slides")
    
    # Add slides from JSON
    for idx, slide_data in enumerate(data["slides"], 1):
        print(f"Adding slide {idx}: {slide_data.get('title', 'Untitled')}")
        
        slide_layout = prs.slide_layouts[slide_layout_index]
        slide = prs.slides.add_slide(slide_layout)
        
        add_slide_content(slide, slide_data)
    
    # Save presentation
    print(f"Saving presentation to: {output_path}")
    prs.save(output_path)
    print(f"Successfully created: {output_path}")
    
    # Print open questions if any
    if "open_questions" in data and data["open_questions"]:
        print("\nOpen Questions:")
        for question in data["open_questions"]:
            print(f"  - {question}")


def main():
    """Main execution function."""
    
    # Configuration
    config = {
        "model_to_use": "Claude",
        "requirements_file": "requirements.docx",
        "questionnaire_file": None,
        "presentation_type": "Proposal",
        "output_json": "proposal_output.json",
        "template_path": os.path.join(os.getcwd(), "template.pptx"),
        "output_pptx": "Generated_Proposal.pptx",
        "slide_layout_index": 1,
        "max_tokens": 4000
    }
    
    print("=" * 60)
    print("PowerPoint Presentation Generator")
    print("=" * 60)
    
    # Step 1: Generate JSON from requirements
    print("\nStep 1: Generating JSON from requirements document...")
    success, result = generate_json_from_requirements(
        model_to_use=config["model_to_use"],
        requirements_file=config["requirements_file"],
        questionnaire_file=config["questionnaire_file"],
        presentation_type=config["presentation_type"],
        output_file_name=config["output_json"],
        max_tokens=config["max_tokens"]
    )
    
    if not success:
        print(f"Failed to generate JSON: {result}")
        return
    
    json_file = result
    
    # Step 2: Generate PowerPoint from JSON
    print("\nStep 2: Generating PowerPoint presentation from JSON...")
    try:
        generate_presentation_from_json(
            json_file_path=os.path.join(os.getcwd(), json_file),
            template_path=config["template_path"],
            output_path=config["output_pptx"],
            slide_layout_index=config["slide_layout_index"]
        )
        
        print("\n" + "=" * 60)
        print("SUCCESS: Presentation generated successfully!")
        print("=" * 60)
        
    except Exception as e:
        print(f"\nError generating PowerPoint: {e}")
        traceback.print_exc()


if __name__ == "__main__":
    main()